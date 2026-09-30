#!/usr/bin/env python3
"""Merge results/*.json into an .xlsx report (uploaded to Drive as a Google Sheet)."""
import glob, json, os, statistics
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import argparse
ap = argparse.ArgumentParser(description=__doc__)
ap.add_argument("--dir", default=os.path.dirname(os.path.abspath(__file__)), help="run dir with results/")
ap.add_argument("--clients", default=None, help="clients JSON (default <dir>/clients.json)")
ap.add_argument("--out", default=None, help="output .xlsx path")
ap.add_argument("--title", default="2026 Appraisal Claims Report")
args = ap.parse_args()
HERE = args.dir
clients = {c["id"]: c for c in json.load(open(args.clients or os.path.join(HERE, "clients.json")))}
results = {}
for p in glob.glob(os.path.join(HERE, "results", "*.json")):
    try:
        r = json.load(open(p))
        rid = r["id"]
        if rid in clients or str(rid) in clients:
            results[rid if rid in clients else str(rid)] = r
    except Exception as e:  # noqa: BLE001
        print("BAD", p, e)

def g(d, *ks):
    for k in ks:
        if not isinstance(d, dict):
            return None
        d = d.get(k)
    return d

def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace("$", "").replace(",", ""))
    except ValueError:
        return None

def appraisal_rows(c):
    return [s for s in c["sheet_rows"] if s.get("type") and "Appraisal" in s["type"]]

MONEY = '"$"#,##0.00'
SCOPE_ITEMS = ["House roof", "Garage roof", "Siding", "Gutters & downspouts", "Windows / screens / wraps", "Interior", "Code upgrade / O&L"]

def scope_text(est):
    sc = (est or {}).get("scope") or []
    return "; ".join(f'{x.get("item")}: {num(x.get("rcv")) or 0:,.0f}' + (f' ({x.get("detail")})' if x.get("detail") else "") for x in sc)

def scope_amt(est, item):
    sc = (est or {}).get("scope") or []
    vals = [num(x.get("rcv")) for x in sc if str(x.get("item", "")).lower() == item.lower()]
    vals = [v for v in vals if v is not None]
    return sum(vals) if vals else None
COLS = [
    # (header, width, fmt, getter)
    ("#", 5, None, lambda c, r: c["id"]),
    ("Client", 26, None, lambda c, r: c["client"]),
    ("Property Address", 36, None, lambda c, r: r.get("address") or c["address"]),
    ("Carrier", 16, None, lambda c, r: r.get("carrier") or c["insurer"]),
    ("Claim #", 16, None, lambda c, r: r.get("claim_number")),
    ("Policy #", 14, None, lambda c, r: r.get("policy_number")),
    ("Date of Loss", 12, None, lambda c, r: r.get("date_of_loss")),
    ("Cause of Loss", 14, None, lambda c, r: r.get("cause_of_loss")),
    ("Deductible", 12, MONEY, lambda c, r: num(r.get("deductible"))),
    ("Appraiser (sheet)", 12, None, lambda c, r: ", ".join(sorted({s["appraiser"] for s in appraisal_rows(c) if s.get("appraiser")}))),
    ("Approval Specialist", 14, None, lambda c, r: ", ".join(sorted({s["specialist"] for s in appraisal_rows(c) if s.get("specialist")}))),
    ("Deal Type (sheet)", 22, None, lambda c, r: "; ".join(s["type"] for s in c["sheet_rows"] if s.get("type"))),
    ("Result Date (sheet)", 12, None, lambda c, r: "; ".join(str(s["result"])[:10] for s in appraisal_rows(c) if s.get("result"))),
    ("Initial Estimate Date", 12, None, lambda c, r: g(r, "initial_estimate", "date")),
    ("Initial Estimate RCV", 14, MONEY, lambda c, r: num(g(r, "initial_estimate", "rcv"))),
    ("Initial Estimate ACV", 14, MONEY, lambda c, r: num(g(r, "initial_estimate", "acv"))),
    ("Initial Net Payment", 14, MONEY, lambda c, r: num(g(r, "initial_estimate", "net_payment"))),
    ("Reinspection Estimate Date", 12, None, lambda c, r: g(r, "reinspection_estimate", "date")),
    ("Reinspection Estimate RCV", 14, MONEY, lambda c, r: num(g(r, "reinspection_estimate", "rcv"))),
    ("Reinspection Estimate ACV", 14, MONEY, lambda c, r: num(g(r, "reinspection_estimate", "acv"))),
    ("Pre-Appraisal Carrier RCV", 14, MONEY, lambda c, r: num(r.get("pre_appraisal_carrier_rcv"))),
    ("Award Date", 12, None, lambda c, r: g(r, "appraisal_award", "date")),
    ("Appraisal Award RCV", 14, MONEY, lambda c, r: num(g(r, "appraisal_award", "total_rcv"))),
    ("Appraisal Award ACV", 14, MONEY, lambda c, r: num(g(r, "appraisal_award", "total_acv"))),
    ("Appraisal Increase", 14, MONEY, lambda c, r: num(r.get("appraisal_increase"))),
    ("Increase (Approvals sheet)", 14, MONEY, lambda c, r: sum(num(s.get("amount")) or 0 for s in appraisal_rows(c)) or None),
    ("Matches Sheet?", 9, None, lambda c, r: {True: "Yes", False: "No"}.get(r.get("sheet_amount_match"), "")),
    ("Award Breakdown", 40, None, lambda c, r: "; ".join(
        f'{l.get("coverage")}: RCV {num(l.get("rcv")) or 0:,.2f} / ACV {num(l.get("acv")) or 0:,.2f}'
        for l in (g(r, "appraisal_award", "lines") or []))),
    ("Award Signatures", 22, None, lambda c, r: g(r, "appraisal_award", "signatures")),
    ("Umpire", 14, None, lambda c, r: g(r, "appraisal_award", "umpire")),
    ("Scope Summary", 50, None, lambda c, r: r.get("scope_summary")),
    ("Initial Estimate Scope", 50, None, lambda c, r: scope_text(r.get("initial_estimate"))),
    ("Reinspection Estimate Scope", 50, None, lambda c, r: scope_text(r.get("reinspection_estimate"))),
    ("Award Scope", 50, None, lambda c, r: scope_text(r.get("appraisal_award"))),
] + [(f"Award {it} RCV", 12, MONEY, (lambda it: lambda c, r: scope_amt(r.get("appraisal_award"), it))(it)) for it in SCOPE_ITEMS] + [
    ("Roof Area (sq ft)", 10, "#,##0", lambda c, r: num(g(r, "property", "eagleview_total_sqft"))),
    ("Roof Squares", 9, "0.00", lambda c, r: num(g(r, "property", "eagleview_squares")) or num(g(r, "property", "carrier_sketch_squares"))),
    ("Pitch", 7, None, lambda c, r: g(r, "property", "predominant_pitch")),
    ("Structures", 34, None, lambda c, r: g(r, "property", "structures")),
    ("Roof / Property Details", 40, None, lambda c, r: r.get("roof_details")),
    ("Other Details", 50, None, lambda c, r: r.get("other_details")),
    ("CC Project ID", 11, None, lambda c, r: r.get("project_id")),
    ("CC Status", 16, None, lambda c, r: r.get("project_status")),
    ("Confidence", 10, None, lambda c, r: r.get("confidence") or "not processed"),
    ("Issues", 60, None, lambda c, r: "; ".join(r.get("issues") or [])),
    ("Files Read", 50, None, lambda c, r: "; ".join(r.get("files_read") or [])),
]

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
FLAG = {"low": PatternFill("solid", fgColor="F8D7DA"), "medium": PatternFill("solid", fgColor="FFF3CD")}

def header(ws, names, widths):
    ws.append(names)
    for i, w in enumerate(widths, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 42

wb = Workbook()
ws = wb.active
ws.title = "Claims"
header(ws, [c[0] for c in COLS], [c[1] for c in COLS])
conf_col = [c[0] for c in COLS].index("Confidence") + 1
for cid in sorted(clients):
    c, r = clients[cid], results.get(cid, {})
    ws.append([f(c, r) for _, _, _, f in COLS])
    row = ws.max_row
    for i, (_, _, fmt, _) in enumerate(COLS, 1):
        if fmt:
            ws.cell(row=row, column=i).number_format = fmt
    fill = FLAG.get(r.get("confidence")) if r else FLAG["low"]
    if fill:
        ws.cell(row=row, column=conf_col).fill = fill
ws.auto_filter.ref = ws.dimensions

# Award line items, one row per coverage line
wl = wb.create_sheet("Award Breakdown")
header(wl, ["#", "Client", "Coverage", "Award RCV", "Award ACV"], [5, 26, 30, 14, 14])
for cid in sorted(results):
    for l in g(results[cid], "appraisal_award", "lines") or []:
        wl.append([cid, clients[cid]["client"], l.get("coverage"), num(l.get("rcv")), num(l.get("acv"))])
        for col in (4, 5):
            wl.cell(row=wl.max_row, column=col).number_format = MONEY

# Carrier estimate history
we = wb.create_sheet("Carrier Estimates")
header(we, ["#", "Client", "Estimate", "Date", "RCV", "ACV", "Net Payment", "File", "Notes"], [5, 26, 22, 12, 14, 14, 14, 36, 50])
for cid in sorted(results):
    r = results[cid]
    ests = [("Initial", r.get("initial_estimate"))] + [("Other carrier estimate", e) for e in r.get("other_carrier_estimates") or []] + [("Reinspection (latest)", r.get("reinspection_estimate"))]
    for label, e in ests:
        if e:
            we.append([cid, clients[cid]["client"], label, e.get("date"), num(e.get("rcv")), num(e.get("acv")), num(e.get("net_payment")), e.get("file"), e.get("notes")])
            for col in (5, 6, 7):
                we.cell(row=we.max_row, column=col).number_format = MONEY

# Scope breakdown, one row per estimate component
wsb = wb.create_sheet("Scope Breakdown")
header(wsb, ["#", "Client", "Estimate", "Date", "Item", "RCV", "Detail"], [5, 26, 22, 12, 26, 14, 70])
for cid in sorted(results):
    r = results[cid]
    ests = [("Initial", r.get("initial_estimate"))] + [("Other carrier estimate", e) for e in r.get("other_carrier_estimates") or []] + [("Reinspection (latest)", r.get("reinspection_estimate")), ("Appraisal award", r.get("appraisal_award"))]
    for label, e in ests:
        for x in (e or {}).get("scope") or []:
            wsb.append([cid, clients[cid]["client"], label, (e or {}).get("date"), x.get("item"), num(x.get("rcv")), x.get("detail")])
            wsb.cell(row=wsb.max_row, column=6).number_format = MONEY

# Summary
done = [results[k] for k in results]
incs = [num(r.get("appraisal_increase")) for r in done if num(r.get("appraisal_increase")) is not None]
awards = [num(g(r, "appraisal_award", "total_rcv")) for r in done if num(g(r, "appraisal_award", "total_rcv")) is not None]
ss = wb.create_sheet("Summary", 0)
ss.column_dimensions["A"].width, ss.column_dimensions["B"].width = 44, 18
rows = [
    (args.title, None),
    ("Source: Approval Department sheet (Data tab, 2026 rows with an appraisal result) + Contractors Cloud claim files", None),
    (None, None),
    ("Claims in scope", len(clients)),
    ("Claims processed", len(done)),
    ("High confidence", sum(r.get("confidence") == "high" for r in done)),
    ("Medium confidence", sum(r.get("confidence") == "medium" for r in done)),
    ("Low confidence", sum(r.get("confidence") == "low" for r in done)),
    ("Increase matches Approvals sheet", sum(r.get("sheet_amount_match") is True for r in done)),
    ("Increase differs from Approvals sheet", sum(r.get("sheet_amount_match") is False for r in done)),
    (None, None),
    ("Claims with an award found", len(awards)),
    ("Total award RCV", sum(awards)),
    ("Total appraisal increase (from files)", sum(incs)),
    ("Average appraisal increase", statistics.mean(incs) if incs else None),
    ("Median appraisal increase", statistics.median(incs) if incs else None),
]
for a, b in rows:
    ss.append([a, b])
ss["A1"].font = Font(bold=True, size=14)
for row in range(13, 17):
    ss.cell(row=row, column=2).number_format = MONEY

out = args.out or os.path.join(HERE, "2026_Appraisal_Claims_Report.xlsx")
wb.save(out)
print(out, "processed", len(done), "of", len(clients))
